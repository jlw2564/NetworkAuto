import os
import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.dirname(__file__))

f = os.path.join(ROOT, '在职人员权限提取结果.xlsx')
g = os.path.join(ROOT, '在职人员-未匹配权限表.xlsx')

exist_f = os.path.isfile(f)
exist_g = os.path.isfile(g)
print('在职人员权限提取结果.xlsx 存在?', exist_f, 'size=', os.path.getsize(f) if exist_f else None)
print('在职人员-未匹配权限表.xlsx 存在?', exist_g, 'size=', os.path.getsize(g) if exist_g else None)

for p in [f, g]:
    if not os.path.isfile(p):
        continue
    wb = load_workbook(p, read_only=True, data_only=True)
    print()
    print('---', os.path.basename(p), 'sheets =', wb.sheetnames)
    wb.close()
    for s in wb.sheetnames:
        df = pd.read_excel(p, sheet_name=s)
        print()
        print('   sheet =', s, 'shape =', df.shape)
        print(df.head(8).to_string())
